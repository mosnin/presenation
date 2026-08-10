"""Export rendered documents: PDF via headless Chromium, DOCX via python-docx."""

from __future__ import annotations

import subprocess
import tempfile
from pathlib import Path

from .theme import Theme


def html_to_pdf(html_text: str, out_path: str | Path, chromium: str) -> Path:
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        "w", suffix=".html", delete=False, dir=out_path.parent
    ) as f:
        f.write(html_text)
        html_path = f.name
    try:
        subprocess.run(
            [
                chromium,
                "--headless=new",
                "--no-sandbox",
                "--disable-gpu",
                "--allow-file-access-from-files",
                "--no-pdf-header-footer",
                f"--print-to-pdf={out_path}",
                f"file://{html_path}",
            ],
            check=True,
            capture_output=True,
            timeout=120,
        )
    finally:
        Path(html_path).unlink(missing_ok=True)
    return out_path


def _hex_to_rgb(hex_color: str) -> tuple[int, int, int]:
    return tuple(int(hex_color[i : i + 2], 16) for i in (1, 3, 5))  # type: ignore[return-value]


def doc_to_docx(doc: dict, theme: Theme, out_path: str | Path) -> Path:
    """Native-DOCX rendering of the same structure. Coarser than the PDF (Word
    styling is style-based, not CSS), but headings, accents, fonts, tables and
    stat rows all carry the template look."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    theme = theme.for_print()  # Word pages are white paper too
    accent = RGBColor(*_hex_to_rgb(theme.accent))
    accent_dark = RGBColor(*_hex_to_rgb(theme.accent_dark))
    ink = RGBColor(*_hex_to_rgb(theme.ink))
    muted = RGBColor(0x6B, 0x72, 0x80)

    d = Document()
    normal = d.styles["Normal"]
    normal.font.name = theme.body_font
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = ink

    title = d.add_paragraph()
    run = title.add_run(doc.get("title", "Document").upper())
    run.font.name = theme.heading_font
    run.font.size = Pt(34)
    run.font.bold = True
    run.font.color.rgb = ink

    if doc.get("subtitle"):
        p = d.add_paragraph()
        run = p.add_run(doc["subtitle"])
        run.font.size = Pt(13)
        run.font.color.rgb = muted

    if doc.get("meta"):
        p = d.add_paragraph()
        run = p.add_run(doc["meta"])
        run.font.size = Pt(9)
        run.font.color.rgb = muted

    for section in doc.get("sections", []):
        h = d.add_heading(level=1)
        run = h.add_run(section.get("heading", "").upper())
        run.font.name = theme.heading_font
        run.font.size = Pt(16)
        run.font.color.rgb = accent_dark

        for block in section.get("blocks", []):
            kind = block.get("type")
            if kind == "paragraph":
                d.add_paragraph(block.get("text", ""))
            elif kind == "bullets":
                for item in block.get("items", []):
                    d.add_paragraph(item, style="List Bullet")
            elif kind == "quote":
                p = d.add_paragraph(style="Intense Quote")
                run = p.add_run(block.get("text", ""))
                run.font.color.rgb = accent_dark
            elif kind == "stats":
                items = block.get("items", [])[:4]
                if not items:
                    continue
                table = d.add_table(rows=2, cols=len(items))
                table.style = "Light Grid Accent 1"
                for i, item in enumerate(items):
                    value_cell = table.cell(0, i).paragraphs[0]
                    value_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = value_cell.add_run(str(item.get("value", "")))
                    run.font.name = theme.heading_font
                    run.font.size = Pt(20)
                    run.font.bold = True
                    run.font.color.rgb = accent
                    label_cell = table.cell(1, i).paragraphs[0]
                    label_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = label_cell.add_run(str(item.get("label", "")).upper())
                    run.font.size = Pt(7.5)
                    run.font.color.rgb = muted
                d.add_paragraph()
            elif kind == "table":
                header = block.get("header", [])
                rows = block.get("rows", [])
                if not header:
                    continue
                table = d.add_table(rows=1 + len(rows), cols=len(header))
                table.style = "Light Grid Accent 1"
                for i, cell_text in enumerate(header):
                    run = table.cell(0, i).paragraphs[0].add_run(str(cell_text))
                    run.font.bold = True
                    run.font.color.rgb = accent_dark
                for r, row in enumerate(rows, start=1):
                    for i, cell_text in enumerate(row[: len(header)]):
                        table.cell(r, i).paragraphs[0].add_run(str(cell_text))
                d.add_paragraph()

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(out_path))
    return out_path
